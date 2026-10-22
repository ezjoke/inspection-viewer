#!/usr/bin/env python3
"""Export an inspection spreadsheet to JSON for the web viewer.

    python3 tools/export_inspection_json.py REPORT.xlsx --out private/inspection.json \
        [--site-name "My house"] [--mask-file private/denylist.txt]

Then assemble a shareable package with tools/build_package.py.

Free-text fields are scrubbed: document reference numbers (seller declaration,
purchase offer) and every regular expression listed in --mask-file are replaced
by "[masqué]". Review the output before sharing it.
"""

import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path
import openpyxl

MASK = "[masqué]"
# Quebec form reference numbers: seller declaration (DV...) and purchase offer (PA...).
BUILTIN_MASKS = [re.compile(r"\b(?:DV|PA)\d{4,}\b")]
MASKS = list(BUILTIN_MASKS)


def load_mask_file(path: Path) -> None:
    for n, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1):
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        try:
            MASKS.append(re.compile(line, re.IGNORECASE))
        except re.error as exc:
            sys.exit(f"{path}:{n}: invalid regular expression: {exc}")


def scrub(value):
    if not isinstance(value, str):
        return value
    for rx in MASKS:
        value = rx.sub(MASK, value)
    return value

COLOR_MAP = {
    "Vert": "vert",
    "Orange": "orange",
    "Rouge": "rouge",
    "Bleu": "bleu",
    "Violet": "violet",
    "Sans symbole": "inconnu",
}

STATUS_COLOR_MAP = {
    "Conforme": "vert",
    "À corriger — préventif": "orange",
    "À corriger — prioritaire": "rouge",
    "Requis — non inspecté": "bleu",
    "Non conforme BNQ": "violet",
    "Information": "inconnu",
}

COLOR_META = {
    "vert": {"label": "Vert", "fill": "#C6EFCE", "textColor": "#1e4620", "badgeClass": "badge-vert"},
    "orange": {"label": "Orange", "fill": "#FFD9A0", "textColor": "#704200", "badgeClass": "badge-orange"},
    "rouge": {"label": "Rouge", "fill": "#FFC7CE", "textColor": "#9C0006", "badgeClass": "badge-rouge"},
    "bleu": {"label": "Bleu", "fill": "#BDD7EE", "textColor": "#002060", "badgeClass": "badge-bleu"},
    "violet": {"label": "Violet", "fill": "#E4C7F0", "textColor": "#371435", "badgeClass": "badge-violet"},
    "inconnu": {"label": "Sans symbole", "fill": "#D9D9D9", "textColor": "#3f3f3f", "badgeClass": "badge-inconnu"},
}

SEVERITY_ORDER = ["rouge", "violet", "bleu", "orange", "vert", "inconnu"]


def slugify(text: str) -> str:
    """Generate clean slug identifier from text."""
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode("utf-8")
    text = re.sub(r"[^\w\s-]", "", text).strip().lower()
    return re.sub(r"[-\s]+", "-", text)


def parse_bullets(text: str) -> list[dict]:
    """Parse multiline text into structured bullet list preserving indent level."""
    if not text or text == "—":
        return []
    lines = text.splitlines()
    bullets = []
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        # Check indentation level
        lspaces = len(line) - len(line.lstrip(" "))
        level = 1 if lspaces >= 2 else 0
        if stripped.startswith("•"):
            item_text = stripped[1:].strip()
        elif stripped.startswith("-"):
            item_text = stripped[1:].strip()
        else:
            item_text = stripped
        bullets.append({
            "text": item_text,
            "level": level,
            "raw": line,
        })
    return bullets


def extract_observations(ws) -> list[dict]:
    """Extract rows from Observations sheet."""
    items = []
    for r in range(2, ws.max_row + 1):
        zone = scrub(ws.cell(r, 1).value)
        if not zone:
            continue
        article = scrub(ws.cell(r, 2).value or "")
        condition = scrub(ws.cell(r, 3).value or "")
        urgence = ws.cell(r, 4).value or "—"
        profession = scrub(ws.cell(r, 5).value or "—")
        couts = ws.cell(r, 6).value or "—"
        statut = ws.cell(r, 7).value or ""
        obs_text = scrub(ws.cell(r, 8).value or "")
        page = ws.cell(r, 9).value

        # Identify color key
        color_key = "inconnu"
        for prefix, ckey in COLOR_MAP.items():
            if condition.startswith(prefix):
                color_key = ckey
                break
        if color_key == "inconnu" and statut in STATUS_COLOR_MAP:
            color_key = STATUS_COLOR_MAP[statut]

        slug = f"{slugify(zone)}__{slugify(article)}"
        items.append({
            "id": slug,
            "row": r,
            "zone": zone,
            "article": article,
            "condition": condition,
            "colorKey": color_key,
            "urgency": urgence,
            "trade": profession,
            "costEstimate": couts,
            "status": statut,
            "rawObservation": obs_text,
            "bullets": parse_bullets(obs_text),
            "page": page,
            "isActionRequired": color_key in ["rouge", "orange", "violet", "bleu"],
            "isExterior": "ext" in slugify(zone) or zone.startswith("Extérieur"),
        })
    return items


def extract_admin_items(ws) -> list[dict]:
    """Extract rows from Informations générales sheet."""
    items = []
    for r in range(2, ws.max_row + 1):
        zone = scrub(ws.cell(r, 1).value)
        if not zone:
            continue
        article = scrub(ws.cell(r, 2).value or "")
        condition = scrub(ws.cell(r, 3).value or "")
        urgence = ws.cell(r, 4).value or "—"
        profession = scrub(ws.cell(r, 5).value or "—")
        couts = ws.cell(r, 6).value or "—"
        statut = ws.cell(r, 7).value or ""
        obs_text = scrub(ws.cell(r, 8).value or "")
        page = ws.cell(r, 9).value

        color_key = "inconnu"
        for prefix, ckey in COLOR_MAP.items():
            if condition.startswith(prefix):
                color_key = ckey
                break

        slug = f"{slugify(zone)}__{slugify(article)}"
        items.append({
            "id": slug,
            "row": r,
            "zone": zone,
            "article": article,
            "condition": condition,
            "colorKey": color_key,
            "urgency": urgence,
            "trade": profession,
            "costEstimate": couts,
            "status": statut,
            "rawObservation": obs_text,
            "bullets": parse_bullets(obs_text),
            "page": page,
        })
    return items


def extract_legend(ws) -> dict:
    """Extract legend and counts from Légende et synthèse sheet."""
    entries = []
    total_articles = 0
    for r in range(2, ws.max_row + 1):
        sym = ws.cell(r, 1).value
        if not sym:
            continue
        if "Total" in str(sym):
            total_articles = ws.cell(r, 2).value
            continue
        cond_text = scrub(ws.cell(r, 2).value or "")
        urg = ws.cell(r, 3).value or "—"
        stat = ws.cell(r, 4).value or ""
        cout = ws.cell(r, 5).value or "—"
        count = ws.cell(r, 6).value or 0

        color_key = sym.lower() if isinstance(sym, str) else "inconnu"
        if color_key not in COLOR_META:
            color_key = "inconnu"

        entries.append({
            "symbol": sym,
            "colorKey": color_key,
            "description": cond_text,
            "urgency": urg,
            "status": stat,
            "baseCost": cout,
            "count": count,
        })
    return {
        "entries": entries,
        "totalArticles": total_articles or sum(e["count"] for e in entries),
    }


def main():
    ap = argparse.ArgumentParser(description="Export an inspection spreadsheet to JSON.")
    ap.add_argument("excel", type=Path, help="Path to the .xlsx report (kept outside the repository)")
    ap.add_argument("--out", type=Path, required=True, help="Output JSON path (e.g. private/inspection.json)")
    ap.add_argument("--site-name", default="Site sans nom")
    ap.add_argument("--mask-file", type=Path, help="Text file, one regular expression per line, to mask")
    ap.add_argument("--expected-items", type=int, help="Fail if the item count differs")
    args = ap.parse_args()

    if not args.excel.exists():
        sys.exit(f"Error: {args.excel} not found.")
    if args.mask_file:
        load_mask_file(args.mask_file)

    print(f"Loading {args.excel.name}...")
    wb = openpyxl.load_workbook(args.excel, data_only=True)

    items = extract_observations(wb["Observations"])
    admin_items = extract_admin_items(wb["Informations générales"])
    legend = extract_legend(wb["Légende et synthèse"])

    status_counts = {}
    color_counts = {}
    for it in items:
        status_counts[it["status"]] = status_counts.get(it["status"], 0) + 1
        color_counts[it["colorKey"]] = color_counts.get(it["colorKey"], 0) + 1

    print("\n--- Summary Verification ---")
    print(f"Total observation items extracted: {len(items)}")
    print(f"Expected in legend: {legend['totalArticles']}")
    if args.expected_items is not None and len(items) != args.expected_items:
        sys.exit(f"Expected {args.expected_items} items, got {len(items)}")

    print("\nCounts per Color Key:")
    for key in SEVERITY_ORDER:
        cnt = color_counts.get(key, 0)
        leg_cnt = next((e["count"] for e in legend["entries"] if e["colorKey"] == key), None)
        print(f"  {key.upper():8s}: {cnt:3d} (legend expects: {leg_cnt})")
        if leg_cnt is not None and cnt != leg_cnt:
            sys.exit(f"Mismatch for {key}: got {cnt} vs legend {leg_cnt}")

    payload = {
        "meta": {"title": args.site_name, "totalItems": len(items), "severityOrder": SEVERITY_ORDER},
        "legend": legend,
        "items": items,
        "adminItems": admin_items,
    }
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nWritten {args.out} ({args.out.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
